"""Write the emailed workbook: readable the moment it opens in Excel.

Columns, in the order you work them: who they are, how old their site is (the
hook), how to reach them (Website is clickable), and the one-liner Message
ready to copy-paste. Rows are sorted oldest domain first — the most compelling
pitch at the top.
"""

# Rating/Reviews stay OUT of the sheet on purpose (same call as Niche: the
# user finds them useless); the pitch still uses them as social proof.
COLUMNS = ["Business Name", "Category", "Since", "Age", "Website", "Phone",
           "Email", "Address", "Message"]

_WRAP_WIDTHS = {"Business Name": 30, "Address": 36, "Message": 70}
_AUTOFIT_MIN = 8
_AUTOFIT_MAX = 40
_AUTOFIT_PAD = 3
_ROW_HEIGHT = 15
_MAX_ROW_LINES = 8


def _row(t: dict, message: str) -> list:
    registered = t.get("registered")
    return [
        t.get("business_name", ""),
        t.get("category", ""),
        registered.year if registered else "",
        f"{int(t.get('age', 0))} yrs",
        t.get("website", ""),
        t.get("phone", ""),
        t.get("email", ""),
        t.get("address", ""),
        message,
    ]


def write(path: str, leads: list) -> None:
    """leads: [(target_dict, message), ...] -> styled xlsx at path."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    # Oldest domain first: the longer they've sat on the site, the stronger
    # the opener reads.
    def _sort_key(lm):
        registered = lm[0].get("registered")
        return (registered.isoformat() if registered else "9999",
                lm[0].get("business_name", "").lower())

    leads = sorted(leads, key=_sort_key)

    wb = Workbook()
    ws = wb.active
    ws.title = "Leads"

    ws.append(COLUMNS)
    head_fill = PatternFill("solid", fgColor="1F2430")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = head_fill
        cell.alignment = Alignment(vertical="center", horizontal="left",
                                   wrap_text=True)
    ws.row_dimensions[1].height = 30

    wrap = Alignment(wrap_text=True, vertical="top")
    top = Alignment(vertical="top")
    longest = {col: len(col) for col in COLUMNS}
    web_col = COLUMNS.index("Website") + 1

    for target, message in leads:
        values = _row(target, message)
        ws.append(values)
        excel_row = ws.max_row
        lines = 1
        for i, (col, value) in enumerate(zip(COLUMNS, values), 1):
            cell = ws.cell(row=excel_row, column=i)
            width = _WRAP_WIDTHS.get(col)
            cell.alignment = wrap if width else top
            text = "" if value is None else str(value)
            if width and text:
                lines = max(lines, -(-len(text) // width))   # ceil division
            elif text:
                longest[col] = max(longest[col], len(text))
        link = values[web_col - 1]
        if link:
            label = target.get("domain") or "open site"
            cell = ws.cell(row=excel_row, column=web_col)
            cell.value = f'=HYPERLINK("{link}","{label}")'
            cell.style = "Hyperlink"
            cell.alignment = top
            longest["Website"] = max(longest["Website"], len(label))
        ws.row_dimensions[excel_row].height = (
            min(lines, _MAX_ROW_LINES) * _ROW_HEIGHT)

    for i, col in enumerate(COLUMNS, 1):
        fixed = _WRAP_WIDTHS.get(col)
        width = fixed if fixed else min(max(longest[col] + _AUTOFIT_PAD,
                                            _AUTOFIT_MIN), _AUTOFIT_MAX)
        ws.column_dimensions[get_column_letter(i)].width = width

    # Freeze header + business name so scrolling right keeps showing who's who.
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{ws.max_row}"
    wb.save(path)
