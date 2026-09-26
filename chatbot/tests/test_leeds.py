"""Tests aimed at the things that actually go wrong.

Every test is offline. Nothing here touches the network, Gemini, or the real
workbook — MASTER/INTENT paths are redirected to a tmpdir.
"""

import os
import shutil
import subprocess
import sys
from datetime import datetime, timedelta, timezone

import pytest

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)

import config                                            # noqa: E402
from core.models import Signal, ensure_utc, utcnow       # noqa: E402
from core.store import Store, content_hash               # noqa: E402
from enrich import contact                               # noqa: E402
from intent import score as scoring                      # noqa: E402


def sig(**kw) -> Signal:
    base = dict(
        source_uid="reddit:t3_x", platform="reddit", source_detail="smallbusiness",
        username="someone", person_key="u/someone@reddit",
        permalink="https://reddit.com/x", post_title="", body="",
        posted_at=utcnow() - timedelta(hours=2))
    base.update(kw)
    return Signal(**base)


# --- the timezone bug -------------------------------------------------------
# This project is global and the window is a hard 48h cut. A naive datetime
# silently shifts it by the machine's UTC offset (+5:30 here) — enough to drop
# half a day of real leads and keep half a day of stale ones.

def test_naive_datetime_is_rejected_not_guessed():
    with pytest.raises(ValueError, match="naive"):
        ensure_utc(datetime(2026, 7, 11, 12, 0, 0))


def test_epoch_and_iso_both_land_in_utc():
    assert ensure_utc(0).tzinfo == timezone.utc
    assert ensure_utc("2026-07-11T12:00:00Z").hour == 12


def test_window_is_a_hard_gate():
    since = utcnow() - timedelta(hours=48)
    old = sig(posted_at=utcnow() - timedelta(hours=49), body="I need a website")
    assert scoring.gate(old, since) == "outside the window"


# --- buyers vs sellers: the dominant failure mode ----------------------------

@pytest.mark.parametrize("body", [
    "I run a web design agency, DM me for a free audit",
    "[FOR HIRE] Experienced developer, check out my portfolio",
    "I'm a freelance designer taking on clients this month",
])
def test_sellers_are_killed(body):
    score, fired, _ = scoring.rule_score(sig(body=body))
    assert score == 0, f"seller survived: {fired}"


def test_retrospective_is_killed():
    # "I built my site" is the single most common false positive.
    score, fired, _ = scoring.rule_score(
        sig(body="I built my website last month and here is what I learned "
                 "about getting a website for a small business"))
    assert score == 0
    assert any("retrospective" in f for f in fired)


def test_practitioner_jargon_is_killed():
    score, _, _ = scoring.rule_score(
        sig(body="Struggling with schema markup and Core Web Vitals on the "
                 "website I need to fix"))
    assert score == 0


def test_real_buyer_scores_high():
    score, fired, tier = scoring.rule_score(sig(
        post_title="Opening a bakery next month and I have no website",
        body="I just registered my business. I have an instagram but no website "
             "yet. How do I get a website made? What should I budget?"))
    assert score >= 70, fired
    assert tier == "stated_problem"


# --- the recall bug ---------------------------------------------------------
# Literal phrase lists miss the phrasings real people use, and the misses are
# the BEST leads. "we need a NEW website" does not contain "need a website".

@pytest.mark.parametrize("body", [
    "We need a new website, our current one is 8 years old",
    "I need a simple 5 page website for a bakery opening next month",
    "Looking for someone who can build our site before we open",
    "We are seeking a freelancer to redesign our online store",
])
def test_natural_phrasings_are_caught(body):
    score, fired, _ = scoring.rule_score(sig(body=body, post_title="help"))
    assert score > 0, f"missed a real buyer: {body!r}"


# --- the gate bug -----------------------------------------------------------
# Structural sources have no author BY CONSTRUCTION. Treating an empty author as
# a deleted account binned every contactable lead in the sheet.

def test_authorless_structural_signals_are_not_dropped_as_bots():
    since = utcnow() - timedelta(hours=48)
    osm = sig(source_uid="osm:node/1", platform="directories", username="",
              person_key="", source_detail="osm:Leeds", signal_type="structural",
              phone="+44 113 000 0000", body="A cafe with a phone and no website")
    assert scoring.gate(osm, since) is None


def test_reddit_signal_with_no_author_IS_dropped():
    since = utcnow() - timedelta(hours=48)
    assert scoring.gate(sig(username="", person_key=""), since) \
        == "bot/deleted author"


def test_supply_side_subreddit_is_dropped():
    since = utcnow() - timedelta(hours=48)
    assert "supply-side" in scoring.gate(
        sig(source_detail="forhire", body="I need a website for my shop"), since)


# --- contact extraction -----------------------------------------------------

def test_junk_emails_are_not_reported_as_contacts():
    assert contact.extract("write to name@example.com")["email"] == ""


def test_real_email_is_found():
    assert contact.extract("Email: owner@bakery.co.uk please")["email"] \
        == "owner@bakery.co.uk"


def test_a_year_is_not_a_phone_number():
    assert contact.extract("we opened in 2019 and grew 20%")["phone"] == ""


def test_country_from_entity_suffix_and_currency():
    assert contact.detect_country("Jones Plumbing Ltd")[0] == "GB"
    assert contact.detect_country("budget is around ₹20,000")[0] == "IN"
    # We must not guess when we cannot tell — a wrong currency ends the pitch.
    assert contact.detect_country("I need a website")[0] == ""


def test_channel_is_public_reply_when_there_is_no_contact():
    assert contact.channel("reddit", "", "") == "public_reply"
    assert contact.channel("jobboards", "", "") == "bid"
    assert contact.channel("directories", "", "+44 113 000") == "phone"


# --- dedup ------------------------------------------------------------------

def test_crossposts_collapse_to_one_content_hash():
    # The same question, reworded and reordered, posted to five subs. The hash is
    # over the sorted set of significant words, so punctuation, capitalisation
    # and word order all fall out.
    a = content_hash("I need a website for my new bakery in Leeds")
    b = content_hash("for my new bakery, in Leeds -- I need a WEBSITE!")
    assert a == b


def test_different_questions_do_not_collide():
    a = content_hash("I need a website for my new bakery in Leeds")
    b = content_hash("my plumbing company is not ranking on Google at all")
    assert a != b


def test_the_same_post_never_resurfaces(tmp_path):
    store = Store(str(tmp_path / "t.db"))
    s = sig()
    assert store.add_signals([s]) == 1
    assert store.add_signals([s]) == 0        # UNIQUE(source_uid) does the work
    store.close()


# --- rate limiting ----------------------------------------------------------

def test_bucket_state_survives_a_new_process(tmp_path):
    """A bucket that resets per run silently doubles your rate when you run the
    tool twice in five minutes. That is how a key gets suspended."""
    from core.ratelimit import RateLimiter
    db = str(tmp_path / "t.db")

    store = Store(db)
    RateLimiter(store.conn).acquire("hackernews")
    before = store.conn.execute(
        "SELECT tokens FROM buckets WHERE key='hackernews'").fetchone()[0]
    store.close()

    store2 = Store(db)                         # a fresh "process"
    after = store2.conn.execute(
        "SELECT tokens FROM buckets WHERE key='hackernews'").fetchone()[0]
    store2.close()
    assert after == before, "bucket did not persist across runs"


# --- draft validation -------------------------------------------------------

def test_draft_that_claims_to_have_seen_a_nonexistent_site_is_rejected():
    from core.models import Lead
    from draft.generate import validate
    lead = Lead(signal=sig(website=""))
    why = validate(
        "I had a look at your website and noticed the homepage is slow, which "
        "is costing you customers every single day of the week right now.", lead)
    assert why and "seen a site" in why


def test_draft_that_invents_a_metric_is_rejected():
    from core.models import Lead
    from draft.generate import validate
    lead = Lead(signal=sig(website="http://x.com", body="my site is slow"))
    why = validate(
        "Your site takes 8 seconds to load, which is well past the point where "
        "most visitors give up and go somewhere else entirely instead.", lead)
    assert why and "metric" in why


def test_a_grounded_draft_passes():
    from core.models import Lead
    from draft.generate import validate
    lead = Lead(signal=sig(website="", body="opening a bakery, no website"))
    ok = validate(
        "For a bakery the thing that moves the needle first is the free Google "
        "Business listing, because that is what shows up when somebody searches "
        "for a bakery near them. A one-page site with your hours and location "
        "covers the rest. Happy to point you at the steps.", lead)
    assert ok is None, ok


# --- the two safety properties ----------------------------------------------

def test_full_mock_run_never_touches_leads_master(tmp_path):
    master = os.path.join(os.path.dirname(ROOT), "leads_master.xlsx")
    if not os.path.exists(master):
        pytest.skip("no leads_master.xlsx on this machine")

    before = os.stat(master)
    digest_before = open(master, "rb").read()

    env = dict(os.environ,
               INTENT_FILE=str(tmp_path / "out.xlsx"),
               LEEDS_DB=str(tmp_path / "t.db"))
    result = subprocess.run(
        [sys.executable, "main.py", "all", "--mock", "--dry-run", "--mock-llm"],
        cwd=ROOT, env=env, capture_output=True, text=True, timeout=180)

    assert result.returncode == 0, result.stderr[-2000:]
    after = os.stat(master)
    assert after.st_mtime == before.st_mtime, "leads_master.xlsx was MODIFIED"
    assert open(master, "rb").read() == digest_before, "leads_master.xlsx changed"
    # A mock run writes the SANDBOX pair, never the real files.
    assert os.path.exists(tmp_path / "out.mock.xlsx")
    assert not os.path.exists(tmp_path / "out.xlsx"), \
        "a mock run wrote the real sheet"
    assert not os.path.exists(tmp_path / "t.db"), \
        "a mock run wrote the real store"


def test_human_edits_survive_a_re_export(tmp_path):
    from openpyxl import load_workbook

    # Mock runs write the sandbox pair, so the sheet under test is .mock.xlsx.
    out = str(tmp_path / "out.mock.xlsx")
    env = dict(os.environ, INTENT_FILE=str(tmp_path / "out.xlsx"),
               LEEDS_DB=str(tmp_path / "t.db"))
    cmd = [sys.executable, "main.py", "all", "--mock", "--dry-run", "--mock-llm"]
    subprocess.run(cmd, cwd=ROOT, env=env, capture_output=True, text=True,
                   timeout=180)

    wb = load_workbook(out)
    ws = wb["Leads"]
    header = [c.value for c in ws[1]]
    ws.cell(row=2, column=header.index("Notes") + 1).value = "DO NOT LOSE ME"
    wb.save(out)

    subprocess.run(cmd, cwd=ROOT, env=env, capture_output=True, text=True,
                   timeout=180)

    ws2 = load_workbook(out)["Leads"]
    notes = [r[header.index("Notes")]
             for r in ws2.iter_rows(min_row=2, values_only=True)]
    assert "DO NOT LOSE ME" in notes, "the export erased a human's notes"


# --- display timezone -------------------------------------------------------
# The sheet is read by a human in India. Storage stays UTC (the window depends on
# it); only the rendering is local.

def test_sheet_renders_local_time_but_storage_stays_utc():
    from datetime import datetime as dt
    from export.excel import _local, DISPLAY_TZ

    noon_utc = dt(2026, 7, 11, 12, 0, 0, tzinfo=timezone.utc)
    assert _local(noon_utc) == "2026-07-11 17:30:00"   # +05:30
    assert DISPLAY_TZ.utcoffset(None) == timedelta(minutes=330)
    # The stored value must NOT have been mutated into local time.
    assert noon_utc.tzinfo == timezone.utc


def test_the_window_is_still_computed_in_utc():
    """If a local timestamp ever leaks into the window it shifts the cut by
    5.5 hours and starts silently dropping live leads."""
    assert utcnow().tzinfo == timezone.utc


def test_the_sheet_carries_no_analyst_columns():
    """The human asked for these to go: Lead ID is an internal key, the two
    timestamps are redundant next to Age (hours), Geo Method is provenance
    trivia, and the rule/LLM score split is debugging detail. One Score."""
    from export.excel import COLUMNS
    for gone in ("Lead ID", "Geo Method", "Score (rules)", "Score (LLM)",
                 "Intent Score"):
        assert gone not in COLUMNS, f"{gone} is back in the sheet"
    assert not any(c.startswith(("Posted At", "Collected At")) for c in COLUMNS)
    for kept in ("Age (hours)", "Score", "WhatsApp", "Permalink"):
        assert kept in COLUMNS, f"{kept} missing from the sheet"


# --- fixtures must not look real --------------------------------------------

def test_mock_permalinks_in_the_SHEET_cannot_open_a_real_post(tmp_path):
    """Reddit ids are base36 and HN ids are plain integers, so an INVENTED
    fixture id is a perfectly valid REAL id: '/comments/aaa3/' opens a genuine
    2005 thread, and HN item 40000001 is somebody's actual comment.

    Sanitising the fixture FILES is not enough — the collectors BUILD the
    permalink from the id, so clean fixtures still produce real URLs. The
    assertion has to be on the thing a human actually clicks: the workbook.
    """
    from openpyxl import load_workbook

    env = dict(os.environ, INTENT_FILE=str(tmp_path / "out.xlsx"),
               LEEDS_DB=str(tmp_path / "t.db"))
    r = subprocess.run(
        [sys.executable, "main.py", "all", "--mock", "--dry-run", "--mock-llm"],
        cwd=ROOT, env=env, capture_output=True, text=True, timeout=180)
    assert r.returncode == 0, r.stderr[-2000:]

    ws = load_workbook(str(tmp_path / "out.mock.xlsx"))["Leads"]
    header = [c.value for c in ws[1]]
    at = header.index("Permalink")
    links = [row[at] for row in ws.iter_rows(min_row=2, values_only=True)]
    assert links, "no leads exported"
    for link in links:
        assert str(link).startswith("https://mock.invalid/"), \
            f"a MOCK lead links to a page that may be real: {link}"


# --- regressions: bugs the suite above did not catch ------------------------

def _stored_lead(store, **kw):
    """A scored lead sitting in a real store, ready to export."""
    from core.models import Lead
    s = sig(**kw)
    store.add_signals([s])
    lead = Lead(signal=s, score_rule=60, score=60)
    store.save_score(lead)
    return lead


def test_human_edits_survive_when_the_runs_tab_was_left_active(tmp_path):
    """THE edit-loss path. Export writes two sheets and invites you to read the
    Runs one. Excel persists whichever tab was on top when you saved, so a human
    who checks Runs and saves leaves activeTab=1 — and _read_human_edits used to
    read wb.active, find no 'Lead ID', conclude "no edits", and overwrite every
    Status/Note in the sheet."""
    from openpyxl import load_workbook
    from export import excel

    path = str(tmp_path / "out.xlsx")
    store = Store(str(tmp_path / "t.db"))
    lead = _stored_lead(store, body="I run a bakery and need a website built")
    excel.export([lead], store, path=path)

    wb = load_workbook(path)
    ws = wb["Leads"]
    header = [c.value for c in ws[1]]
    ws.cell(row=2, column=header.index("Notes") + 1).value = "CALLED HIM"
    ws.cell(row=2, column=header.index("Do Not Contact") + 1).value = "yes"
    wb.active = wb.index(wb["Runs"])          # exactly what Excel saves
    wb.save(path)

    excel.export([lead], store, path=path)    # the next run

    ws2 = load_workbook(path)["Leads"]
    h2 = [c.value for c in ws2[1]]
    assert ws2.cell(row=2, column=h2.index("Notes") + 1).value == "CALLED HIM"
    assert ws2.cell(row=2, column=h2.index("Do Not Contact") + 1).value == "yes"
    assert store.human_fields(lead.signal.source_uid)["notes"] == "CALLED HIM"
    store.close()


def test_export_refuses_to_write_the_siblings_workbook(tmp_path):
    from export import excel
    store = Store(str(tmp_path / "t.db"))
    with pytest.raises(RuntimeError, match="leads_master"):
        excel.export([], store, path=str(tmp_path / "leads_master.xlsx"))
    store.close()


def test_control_characters_in_a_post_do_not_kill_the_export(tmp_path):
    """openpyxl raises IllegalCharacterError on \\x0b/\\x0c/\\x07, which real
    bodies contain (pasted terminal output). Uncaught, it destroyed the run
    AFTER collect/score/enrich/draft had already done all the work."""
    from export import excel
    store = Store(str(tmp_path / "t.db"))
    lead = _stored_lead(
        store,
        post_title="Need a website \x07for my cafe",
        body="Error I get:\x0b Traceback\x0c ... and I need a website \U0001f600")
    path = excel.export([lead], store, path=str(tmp_path / "out.xlsx"))
    assert os.path.exists(path)
    store.close()


@pytest.mark.parametrize("body", [
    # "for free" is a substring of "for freelancers" -> hard-killed as no_budget
    "I run a plumbing company and we do not have a website yet. I am looking "
    "for freelancers who can build one. What is a realistic budget?",
    # "react" is a substring of "reaction" -> hard-killed as practitioner jargon
    "I own a small bakery. The reaction to our new menu has been great but we "
    "have no website yet and I want a website built. How much does it cost?",
])
def test_a_buyer_is_not_killed_by_a_substring_of_another_word(body):
    score, fired, _ = scoring.rule_score(sig(body=body, post_title="help"))
    assert score > 0, f"a real buyer was hard-killed: {fired}"


def test_phrase_banks_still_match_whole_phrases():
    from intent import keywords as kw
    assert kw.hits("check out my portfolio: here", kw.SELLER)
    assert kw.hits("[FOR HIRE] senior dev".lower(), kw.SELLER)
    assert kw.hits("my budget is $500", kw.BUDGET)
    assert not kw.hits("we are looking for freelancers", kw.NO_BUDGET)


def test_no_gemini_key_never_fabricates_a_verdict_or_a_draft(tmp_path):
    """No key used to silently become MOCK MODE: every borderline lead got an
    invented verdict (business_type 'bakery', region 'Leeds', country 'GB', a
    made-up stated_problem) and an invented opener, written to the DB and the
    sheet as if a model had produced them."""
    from core.gemini import Gemini, GeminiError
    from draft import generate as drafting

    os.environ.pop("GEMINI_API_KEY", None)
    os.environ.pop("GEMINI_MOCK", None)
    gemini = Gemini(store=None, mock=False)
    assert gemini.enabled is False and gemini.mock is False
    with pytest.raises(GeminiError):
        gemini.structured("sys", "user", {})

    store = Store(str(tmp_path / "t.db"))
    s = sig(post_title="Do I need a website for a market stall?",
            body="I sell handmade candles at a market stall. I only have an "
                 "instagram and I am not sure a website is worth it for me.")
    store.add_signals([s])
    since = utcnow() - timedelta(hours=48)
    leads = scoring.score_all([s], since, gemini, store)
    assert leads, "a missing key must not empty the sheet — rule scores stand"
    v = leads[0].verdict
    assert (v.business_type, v.region, v.country, v.stated_problem) == \
        ("", "", "", ""), f"fabricated verdict: {v}"
    assert leads[0].score_llm == 0

    assert drafting.draft_all(leads, gemini, store) == 0
    opener, status = store.conn.execute(
        "SELECT opener, draft_status FROM signals").fetchone()
    assert opener == ""
    assert "needs_manual" in status
    store.close()


def test_acquire_skips_instead_of_sleeping_for_a_day(tmp_path, monkeypatch):
    """newdomains is (capacity 1, 1 token/86400s): the SECOND feed fetch of a
    normal run asked acquire() to sleep 86400s. The run just hung for a day."""
    from core import ratelimit

    slept: list[float] = []
    monkeypatch.setattr(ratelimit.time, "sleep", lambda s: slept.append(s))
    store = Store(str(tmp_path / "t.db"))
    limiter = ratelimit.RateLimiter(store.conn)

    assert limiter.acquire("newdomains") is True     # day 1 of NRD_DAYS_BACK
    assert limiter.acquire("newdomains") is True     # day 2 — used to sleep 24h
    assert limiter.acquire("newdomains") is False    # daily quota (2) is spent
    assert not slept, f"acquire slept {sum(slept)}s"
    store.close()


def test_regeneration_actually_re_asks_the_model(tmp_path):
    """gemini.structured() caches on hash(system+user). Re-sending the identical
    prompt for a rejected batch replayed the SAME rejected draft out of the
    cache, so REGEN_ROUNDS never re-asked the model once."""
    from core.gemini import Gemini
    from draft import generate as drafting

    class Flaky(Gemini):
        def __init__(self, store):
            super().__init__(store=store, mock=True)
            self.calls = 0

        def _mock(self, user):
            self.calls += 1
            uid = [l[3:].strip() for l in user.splitlines()
                   if l.startswith("ID:")][0]
            if self.calls == 1:      # ungrounded: claims to have seen a site
                return [{"id": uid,
                         "opener": "I had a look at your website and it is slow."}]
            return [{"id": uid, "opener": (
                "For a bakery the first thing worth doing is claiming the free "
                "Google Business listing, because that is what shows up when "
                "somebody searches for a bakery near them. A one page site with "
                "your hours and location covers the rest. Happy to point you at "
                "the steps if that is useful.")}]

    store = Store(str(tmp_path / "t.db"))
    lead = _stored_lead(store, body="opening a bakery, no website yet",
                        contact_channel="public_reply")
    gemini = Flaky(store)
    assert drafting.draft_all([lead], gemini, store) == 1
    assert gemini.calls == 2, "the retry was served from the LLM cache"
    assert store.conn.execute(
        "SELECT draft_status FROM signals").fetchone()[0] == "ok"
    store.close()


def test_a_collector_that_crashes_is_not_recorded_as_a_clean_empty_run():
    """The Runs sheet is what you read when yield drops. A collector that threw
    used to land there as 0 fetched / 0 errors — identical to a quiet day."""
    from collectors.base import Collector

    class Broken(Collector):
        name = "broken"

        def fetch(self, since):
            raise RuntimeError("the API changed shape")

    c = Broken(None, None, None)
    assert c.run(utcnow() - timedelta(hours=48)) == []
    assert c.errors == 1


# --- requirements completeness ----------------------------------------------

def test_every_third_party_import_is_declared():
    """/scraper imports openpyxl and never declares it, so a clean install
    produces a run that reports success and writes zero leads."""
    declared = set()
    with open(os.path.join(ROOT, "requirements.txt"), encoding="utf-8") as fh:
        for line in fh:
            line = line.strip()
            if line and not line.startswith("#"):
                declared.add(line.split(">=")[0].split("==")[0].strip().lower())

    for module in ("requests", "pydantic", "typer", "openpyxl", "dnspython"):
        assert module in declared, f"{module} is imported but not in requirements"


# --- over-broad kill phrases -------------------------------------------------
# A hard-kill is final: score 0, no LLM, no appeal. So a kill phrase that matches
# ordinary business English destroys real buyers invisibly.

@pytest.mark.parametrize("body", [
    # A bare "we offer" killed this. Every trades business describes itself this way.
    "We are a landscaping company. We offer lawn care and snow clearing. "
    "We don't have a website and want to hire a web designer to build one.",
    # A bare "learning" killed this. They are not a student — they are in pain.
    "I own a salon and I am learning the hard way that we need a website. "
    "Customers keep asking for one and I have no idea where to start.",
])
def test_ordinary_business_english_is_not_mistaken_for_a_seller_or_student(body):
    score, fired, _ = scoring.rule_score(sig(body=body, post_title="help please"))
    assert score > 0, f"a real buyer was hard-killed: {fired}"


def test_an_actual_seller_is_still_killed():
    for body in ("We offer web design and SEO services for small businesses.",
                 "I'm learning to code and built my first portfolio project."):
        score, _, _ = scoring.rule_score(sig(body=body))
        assert score == 0, f"seller/student survived: {body!r}"


# --- non-English leads were invisible ---------------------------------------
# Every phrase bank is English, so a non-English buyer scored 0 and score_all
# dropped it BEFORE Gemini saw it. Not rejected -- invisible.

def test_language_detector():
    from intent.keywords import looks_english
    assert looks_english("I need a website for my new bakery, can anyone help?")
    assert looks_english("thanks!")            # too short to judge -> assume EN
    assert not looks_english(
        "Necesito una pagina web para mi negocio nuevo, alguien puede ayudarme")
    assert not looks_english(
        "मुझे अपने व्यव"
        "साय के लिए एक "
        "वेबसाइट चाहिए")


def test_non_english_buyer_reaches_the_classifier_instead_of_the_bin(tmp_path):
    """The whole point: a Spanish buyer scores 0 on English banks. It must be
    handed to Gemini, not silently deleted."""
    from core.gemini import Gemini
    from core.store import Store as S

    store = S(str(tmp_path / "t.db"))
    gemini = Gemini(store=store, mock=True)
    since = utcnow() - timedelta(hours=48)

    spanish = sig(
        source_uid="reddit:t3_es1",
        post_title="Necesito ayuda con mi negocio",
        body="Acabo de abrir una panaderia y no tengo pagina web. "
             "Busco alguien que pueda crear un sitio para mi negocio. "
             "Cuanto cuesta normalmente?")

    # It scores 0 on the English rules -- that is the premise.
    rule, fired, _ = scoring.rule_score(spanish)
    assert rule == 0, "premise broken: the Spanish post matched an English bank"
    assert not any(f.startswith("kill:") for f in fired), "should not be killed"

    # The claim under test is that it REACHED the classifier, not that the
    # classifier liked it. Whether it survives depends on the verdict (the mock
    # flips a deterministic coin per id) and that is the classifier's job, not
    # the router's. What must never happen again is it being binned UNSEEN.
    scoring.score_all([spanish], since, gemini, store)
    seen = store.conn.execute("select count(*) from llm_cache").fetchone()[0]
    assert seen > 0, "the Spanish buyer never reached Gemini -- dropped unseen"
    store.close()


def test_a_non_english_SELLER_is_still_killed_for_free(tmp_path):
    """The escape hatch must not become a hole. A hard-kill still ends it."""
    spanish_seller = sig(
        body="Somos una agencia. We offer web design y SEO. DM me para "
             "una consulta gratis sobre su sitio web.")
    rule, fired, _ = scoring.rule_score(spanish_seller)
    assert rule == 0
    assert any(f.startswith("kill:") for f in fired), fired


def test_foreign_routing_is_capped(tmp_path):
    """Uncapped, this would hand the whole non-English half of Reddit to a
    throttled free-tier Gemini."""
    import config as C
    assert C.MAX_FOREIGN_TO_CLASSIFY > 0
    assert C.MAX_FOREIGN_TO_CLASSIFY <= 500


# --- WhatsApp links -----------------------------------------------------------
# The user is in India. A UK landline in the Phone column is a number he cannot
# dial; the WhatsApp column is the actionable form. A WRONG normalisation
# messages a stranger, so ambiguity must return "" and never a guess.

@pytest.mark.parametrize("phone,country,expected", [
    ("+44 113 496 0000", "", "441134960000"),      # OSM norm: already E.164
    ("0113 496 0000", "GB", "441134960000"),       # UK national form
    ("+91 98765 43210", "", "919876543210"),
    ("9876543210", "IN", "919876543210"),          # bare Indian mobile
    ("(512) 555-0133", "US", "15125550133"),
    ("0113 496 0000", "", ""),                     # national form, no country
    ("12345", "GB", ""),                           # not a phone number
    ("", "GB", ""),
])
def test_to_e164(phone, country, expected):
    assert contact.to_e164(phone, country) == expected


def test_wa_link_is_blank_rather_than_wrong():
    assert contact.wa_link("+44 113 496 0000") == "https://wa.me/441134960000"
    assert contact.wa_link("0113 496 0000", "") == ""


def test_the_sheet_renders_whatsapp_as_a_link(tmp_path):
    from openpyxl import load_workbook
    from export import excel

    store = Store(str(tmp_path / "t.db"))
    lead = _stored_lead(store, source_uid="osm:node/1", platform="directories",
                        username="", person_key="", signal_type="structural",
                        phone="+44 113 496 0000", country="GB",
                        body="A cafe with a phone and no website")
    lead.verdict.service_wanted = "web_design"
    path = excel.export([lead], store, path=str(tmp_path / "out.xlsx"))

    ws = load_workbook(path)["Leads"]
    header = [c.value for c in ws[1]]
    assert ws.cell(row=2, column=header.index("WhatsApp") + 1).value \
        == "https://wa.me/441134960000"
    store.close()


# --- Service Wanted: no "none" rows in the sheet -------------------------------

def test_service_wanted_none_is_left_out_of_the_sheet():
    from core.models import Lead, Verdict
    from export.excel import exportable
    keep = Lead(signal=sig(), verdict=Verdict(service_wanted="web_design"))
    drop = Lead(signal=sig(source_uid="reddit:t3_y"))     # default "none"
    assert exportable([keep, drop]) == [keep]


def test_rule_fallback_names_the_service_when_gemini_cannot(tmp_path):
    """Only Gemini sets service_wanted, and the sheet now drops "none" rows —
    so without a rule fallback, a model outage would blank the whole sheet
    while looking like a clean run."""
    from core.gemini import Gemini

    os.environ.pop("GEMINI_API_KEY", None)
    gemini = Gemini(store=None, mock=False)
    assert gemini.enabled is False

    store = Store(str(tmp_path / "t.db"))
    since = utcnow() - timedelta(hours=48)
    web = sig(source_uid="reddit:t3_w", post_title="help",
              body="Opening a bakery and I need a website built, what budget?")
    seo = sig(source_uid="reddit:t3_s", post_title="help",
              body="My shop is not showing up on google, I need seo help")
    osm = sig(source_uid="osm:node/9", platform="directories", username="",
              person_key="", signal_type="structural", phone="+44 113 000 0000",
              body="A cafe with a phone and no website")
    store.add_signals([web, seo, osm])

    leads = {l.signal.source_uid: l
             for l in scoring.score_all([web, seo, osm], since, gemini, store)}
    assert leads["reddit:t3_w"].verdict.service_wanted == "web_design"
    assert leads["reddit:t3_s"].verdict.service_wanted == "seo"
    assert leads["osm:node/9"].verdict.service_wanted == "web_design"
    store.close()


# --- OSM: "no website tag" must not mean "has a website elsewhere" ------------

def test_osm_business_with_a_brand_or_hidden_website_tag_is_dropped():
    from collectors.directories import DirectoriesCollector

    def element(extra):
        return {"type": "node", "id": 1,
                "tags": {"name": "Greggs", "phone": "+44 113 000 0000",
                         "shop": "bakery", **extra}}

    for extra in ({"brand": "Greggs"}, {"contact:website": "https://x.com"},
                  {"url": "https://x.com"}, {"wikidata": "Q123"}):
        assert DirectoriesCollector._element(element(extra), "Leeds") is None, \
            f"established business survived the tag screen: {extra}"
    kept = DirectoriesCollector._element(element({}), "Leeds")
    assert kept is not None
    assert kept.country == "GB"       # AREA_COUNTRY feeds the WhatsApp link


def test_a_social_page_that_links_to_a_real_site_drops_the_lead(tmp_path):
    from collectors.directories import DirectoriesCollector

    class FakeResp:
        status_code = 200
        text = ('<html><meta property="og:url" content='
                '"https://www.facebook.com/thecafe"/>'
                '<a href="https://www.thecafeleeds.co.uk">Website</a></html>')

    class FakeScraper:
        def get(self, url, bucket="web", **kw):
            return FakeResp()

    class BlockedScraper:
        def get(self, url, bucket="web", **kw):
            return None                       # robots.txt said no / login wall

    store = Store(str(tmp_path / "t.db"))
    c = DirectoriesCollector(None, FakeScraper(), store)
    s = sig(source_uid="osm:node/2", platform="directories", username="",
            person_key="", signal_type="structural", phone="+44 113 000 0000",
            raw={"tags": {"contact:facebook": "thecafe"}})
    assert c._website_via_links(s) == "https://www.thecafeleeds.co.uk"

    # A BLOCKED fetch is not evidence of a website — the lead survives.
    c2 = DirectoriesCollector(None, BlockedScraper(), store)
    assert c2._website_via_links(s) == ""
    store.close()


# --- the merge key survived the Lead ID removal --------------------------------

def test_edits_in_an_old_lead_id_sheet_still_survive(tmp_path):
    """The sheet used to key rows by Lead ID; it keys by Permalink now. The
    FIRST export after that change reads a sheet in the OLD format — losing
    the human's notes there would be the exact bug this file exists to stop."""
    from openpyxl import Workbook, load_workbook
    from export import excel

    store = Store(str(tmp_path / "t.db"))
    lead = _stored_lead(store, body="I run a bakery and need a website built")
    lead.verdict.service_wanted = "web_design"
    path = str(tmp_path / "out.xlsx")

    wb = Workbook()                            # a sheet in the OLD format
    ws = wb.active
    ws.title = "Leads"
    ws.append(["Lead ID", "Post Title", "Status", "Contacted",
               "Do Not Contact", "Notes"])
    ws.append([lead.signal.source_uid, "old row", "", "", "", "OLD NOTE"])
    wb.save(path)

    excel.export([lead], store, path=path)     # first new-format export

    ws2 = load_workbook(path)["Leads"]
    h2 = [c.value for c in ws2[1]]
    assert "Lead ID" not in h2
    assert ws2.cell(row=2, column=h2.index("Notes") + 1).value == "OLD NOTE"
    store.close()


# --- geography ---------------------------------------------------------------

def test_the_directory_sweep_is_india_only():
    """An OSM lead is a cold call to a shop with a phone and no website, so a
    city we cannot ring is not a lead. This used to assert the opposite ("sweep
    more than the UK"); collecting the UK and US filled the sheet with rows that
    could never be worked. Every swept city must be reachable."""
    import config as C
    assert len(C.OVERPASS_AREAS) >= 8, "keep enough cities to fill a sheet"
    for city in ("Delhi", "Gurugram", "Bengaluru"):
        assert city in C.OVERPASS_AREAS, f"{city} missing from OVERPASS_AREAS"
    foreign = [c for c in C.OVERPASS_AREAS if C.AREA_COUNTRY.get(c) != "IN"]
    assert not foreign, f"only India is contactable; drop {foreign}"


def test_areas_flag_overrides_the_default_city_list(tmp_path):
    from collectors.directories import DirectoriesCollector
    from core.store import Store as S
    store = S(str(tmp_path / "t.db"))
    c = DirectoriesCollector(None, None, store, areas=["Delhi", "Gurugram"])
    assert c.areas == ["Delhi", "Gurugram"]
    d = DirectoriesCollector(None, None, store)
    assert len(d.areas) >= 8      # falls back to the full list
    store.close()
