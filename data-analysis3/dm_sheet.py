"""The DM sheet: the 10 best Instagram-reachable businesses, links you can click.

One row per business, two clickable links: the public profile (vet them in ten
seconds) and the ig.me deep link that opens their DM thread directly. No
Rating/Reviews/state columns — those stay in the master file; this sheet is a
work list, and each row is a DM to send today.

Handles come from the business's own Google Maps listing, i.e. the account the
owner chose to publish — in practice a public business profile. ig.me works
either way: a private account just receives the DM as a request.
"""

COLUMNS = ["Business Name", "Instagram", "DM", "Profile", "Category", "Address"]

_WRAP_WIDTHS = {"Business Name": 30, "Address": 36}
_AUTOFIT_MIN = 8
_AUTOFIT_MAX = 40
_AUTOFIT_PAD = 3


def handle(target: dict) -> str:
    return str(target.get("instagram") or "").strip().lstrip("@")


def dm_link(target: dict) -> str:
    h = handle(target)
    return f"https://ig.me/m/{h}" if h else ""


def profile_link(target: dict) -> str:
    h = handle(target)
    return f"https://www.instagram.com/{h}/" if h else ""


def write(path: str, targets: list) -> None:
    """targets: sourcer dicts (instagram non-empty) -> styled xlsx at path."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "DM Leads"

    ws.append(COLUMNS)
    head_fill = PatternFill("solid", fgColor="1F2430")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = head_fill
        cell.alignment = Alignment(vertical="center", horizontal="left",
                                   wrap_text=True)
    ws.row_dimensions[1].height = 30

    wrap = Alignment(wrap_text=True, vertical="top")
    top = Alignment(vertical="top")
    longest = {col: len(col) for col in COLUMNS}

    for t in targets:
        h = handle(t)
        values = [t.get("business_name", ""), f"@{h}" if h else "",
                  "Open DM", "Profile",
                  t.get("category", ""), t.get("address", "")]
        ws.append(values)
        row = ws.max_row
        lines = 1
        for i, (col, value) in enumerate(zip(COLUMNS, values), 1):
            cell = ws.cell(row=row, column=i)
            width = _WRAP_WIDTHS.get(col)
            cell.alignment = wrap if width else top
            text = str(value or "")
            if width and text:
                lines = max(lines, -(-len(text) // width))
            elif text:
                longest[col] = max(longest[col], len(text))
        for col, url, label in (("DM", dm_link(t), "Open DM"),
                                ("Profile", profile_link(t), "Profile")):
            cell = ws.cell(row=row, column=COLUMNS.index(col) + 1)
            if url:
                cell.value = f'=HYPERLINK("{url}","{label}")'
                cell.style = "Hyperlink"
                cell.alignment = top
        ws.row_dimensions[row].height = lines * 15

    for i, col in enumerate(COLUMNS, 1):
        fixed = _WRAP_WIDTHS.get(col)
        width = fixed if fixed else min(max(longest[col] + _AUTOFIT_PAD,
                                            _AUTOFIT_MIN), _AUTOFIT_MAX)
        ws.column_dimensions[get_column_letter(i)].width = width

    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{ws.max_row}"
    wb.save(path)
